# -*- coding: utf-8 -*-
"""개업공인중개사 현황: 거시 개폐업 추세 모니터링 + B2B 영업 타겟팅 DB 추출을 한 화면에서."""
import io

import pandas as pd
import plotly.graph_objects as go
import streamlit as st
import streamlit.components.v1 as components
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import real_stats
import region_utils as ru
from theme import CARD, GREEN, GREEN_DEEP, INK, LINE, MUTED, ROW_HOVER, TABLE_LINE, VERMILION, ACCENT_SOFT, BLUE

_NET_UP_FILL = PatternFill(start_color="E3F0EA", end_color="E3F0EA", fill_type="solid")
_NET_DOWN_FILL = PatternFill(start_color="FBEAEA", end_color="FBEAEA", fill_type="solid")

# 지역/시작/종료/빠른선택을 한 줄에 컴팩트하게 — select box 폭을 줄이고, pill 패딩을 줄이고,
# 줄바꿈 없이 한 줄로 유지한다.
_FILTER_CSS = f"""
<style>
.st-key-br_filter_row [data-testid="stHorizontalBlock"] {{
    flex-wrap: nowrap !important; align-items: center !important; gap: 4px !important;
}}
.st-key-br_filter_row [data-testid="stColumn"] {{ min-width: 0 !important; width: auto !important; flex: 0 0 auto !important; }}
.st-key-br_filter_row [data-testid="stSelectbox"] {{ min-width: 84px !important; max-width: 120px !important; }}
.st-key-br_filter_row [data-testid="stSelectbox"] [role="group"] {{ min-height: 32px !important; }}
.st-key-br_preset_row div[role="radiogroup"] {{ flex-wrap: nowrap !important; gap: 3px !important; }}
.st-key-br_preset_row button[data-variant="segmented_control"] {{
    padding: 4px 10px !important; font-size: .78rem !important; white-space: nowrap !important;
    border-radius: 8px !important; min-height: unset !important;
}}
</style>
"""


def _sum_or_none(series: pd.Series):
    return series.sum(min_count=1)


def _region_label(region: str) -> str:
    if region in ru.SHORT_NAME:
        return ru.SHORT_NAME[region]
    parts = region.split(" ", 1)
    return parts[1] if len(parts) > 1 else region


def _sync_table_region():
    picked = st.session_state.get("br_region")
    if picked:
        st.session_state["br_table_region"] = picked


def _reset_table_region():
    st.session_state["br_table_region"] = "전국"


def _apply_preset(all_months: list, latest_month: str):
    preset = st.session_state.get("br_preset")
    if preset is None:
        return
    year_start = f"{latest_month[:4]}-01"
    year_start = year_start if year_start in all_months else all_months[0]
    mapping = {
        "최근 6개월": (all_months[max(0, len(all_months) - 6)], latest_month),
        "올해": (year_start, latest_month),
        "3개년": (all_months[max(0, len(all_months) - 36)], latest_month),
        "5개년": (all_months[max(0, len(all_months) - 60)], latest_month),
        "전체": (all_months[0], latest_month),
    }
    start, end = mapping[preset]
    st.session_state["br_start"] = start
    st.session_state["br_end"] = end


def _resolve_table_scope(table_region_short: str, base_df: pd.DataFrame, regions_full: list, period_months: list):
    if table_region_short == "전국":
        district_rows = real_stats.load_all_district_monthly_stats()
        d_df = pd.DataFrame(district_rows)
        if not d_df.empty:
            d_df = d_df[d_df["월"].isin(period_months)]
        if d_df.empty:
            return base_df, regions_full
        districts_by_sido = {}
        for district in d_df["지역"].unique():
            districts_by_sido.setdefault(ru.normalize_sido(district), []).append(district)
        ordered_regions = []
        for sido_full in regions_full:
            ordered_regions.append(sido_full)
            ordered_regions.extend(sorted(districts_by_sido.get(sido_full, [])))
        combined = pd.concat([base_df, d_df], ignore_index=True)
        return combined, ordered_regions
    sido_full = ru.FULL_NAME[table_region_short]
    district_rows = real_stats.load_district_monthly_stats(sido_full)
    if district_rows:
        d_df = pd.DataFrame(district_rows)
        d_df = d_df[d_df["월"].isin(period_months)]
        if not d_df.empty:
            return d_df, sorted(d_df["지역"].unique())
    return base_df, [sido_full]


def _net_cell(opened, closed) -> str:
    if pd.isna(opened) or pd.isna(closed):
        return '<td class="bf">–</td>'
    net = int(opened) - int(closed)
    cls = "net-up" if net >= 0 else "net-down"
    return f'<td class="{cls}">{net:+,}</td>'


def _row_cells(df: pd.DataFrame, region: str, months_desc: list) -> str:
    sub = df[df["지역"] == region].set_index("월")
    cells = []
    for m in months_desc:
        if m in sub.index:
            r = sub.loc[m]
            cls = ' class="bf"' if r["출처"] == "backfill" else ""
            o, c, live = r["개업"], r["폐업"], r.get("영업중")
            o_txt = "–" if pd.isna(o) else f"{int(o):,}"
            c_txt = "–" if pd.isna(c) else f"{int(c):,}"
            live_txt = "–" if pd.isna(live) else f"{int(live):,}"
            cells.append(f"<td{cls}>{o_txt}</td><td{cls}>{c_txt}</td>{_net_cell(o, c)}<td{cls}>{live_txt}</td>")
        else:
            cells.append('<td class="bf">–</td><td class="bf">–</td><td class="bf">–</td><td class="bf">–</td>')
    return "".join(cells)


def _build_table_html(df: pd.DataFrame, regions: list, months_desc: list) -> str:
    rows_html = []
    for region in regions:
        label = _region_label(region)
        row_cls = ' style="font-weight:700"' if region == "전국" else ""
        rows_html.append(f'<tr{row_cls}><td class="region">{label}</td>{_row_cells(df, region, months_desc)}</tr>')

    header_months = "".join(f'<th class="mo" colspan="4">{m}</th>' for m in months_desc)
    header_sub = '<th class="region"></th>' + '<th class="oc-o">개업</th><th class="oc-c">폐업</th><th>순증감</th><th>영업중(Live)</th>' * len(months_desc)

    return f"""
    <div class="ov-table-scroll">
    <table class="ov-table">
        <thead><tr><th class="region"></th>{header_months}</tr><tr>{header_sub}</tr></thead>
        <tbody>{''.join(rows_html)}</tbody>
    </table>
    </div>
    """


def _build_nationwide_tree_component(
    table_df: pd.DataFrame, regions_full: list, regions_short: list, period_months: list, months_desc: list,
) -> str:
    """전국 상세표: 시/도 행을 클릭하면 바로 아래로 그 시/도의 시군구 행이 펼쳐진다."""
    rows_html = [f'<tr style="font-weight:700"><td class="region">전국</td>{_row_cells(table_df, "전국", months_desc)}</tr>']
    for region_short in regions_short:
        if region_short == "전국":
            continue
        sido_full = ru.FULL_NAME[region_short]
        district_rows = real_stats.load_district_monthly_stats(sido_full)
        d_df = pd.DataFrame(district_rows)
        if not d_df.empty:
            d_df = d_df[d_df["월"].isin(period_months)]
        districts = sorted(d_df["지역"].unique()) if not d_df.empty else []

        if districts:
            toggle_attrs = f' class="region-toggle" data-toggle="{region_short}"'
            chev = '<span class="chev">▸</span> '
        else:
            toggle_attrs = ""
            chev = ""
        rows_html.append(
            f'<tr{toggle_attrs} data-open="0">'
            f'<td class="region">{chev}{region_short}</td>{_row_cells(table_df, sido_full, months_desc)}</tr>'
        )
        for district in districts:
            label = _region_label(district)
            rows_html.append(
                f'<tr data-group="{region_short}" class="district-row" style="display:none">'
                f'<td class="region district-cell">{label}</td>{_row_cells(d_df, district, months_desc)}</tr>'
            )

    header_months = "".join(f'<th class="mo" colspan="4">{m}</th>' for m in months_desc)
    header_sub = '<th class="region"></th>' + '<th class="oc-o">개업</th><th class="oc-c">폐업</th><th>순증감</th><th>영업중(Live)</th>' * len(months_desc)

    region_col_width = 130
    data_col_width = 62
    colgroup_html = f'<col style="width:{region_col_width}px">' + (
        f'<col style="width:{data_col_width}px"><col style="width:{data_col_width}px"><col style="width:{data_col_width}px"><col style="width:{data_col_width}px">' * len(months_desc)
    )
    table_width = region_col_width + data_col_width * 4 * len(months_desc)

    return f"""
    <!doctype html>
    <html><head><meta charset="utf-8">
    <style>
    @import url('https://cdn.jsdelivr.net/gh/orioncactus/pretendard@v1.3.9/dist/web/variable/pretendardvariable-dynamic-subset.min.css');
    * {{ box-sizing:border-box; font-family:"Pretendard Variable",Pretendard,-apple-system,"Malgun Gothic",sans-serif; }}
    body {{ margin:0; color:{INK}; font-variant-numeric:tabular-nums; }}
    .scroll {{ overflow-x:auto; }}
    table {{ table-layout:fixed; border-collapse:collapse; font-size:.82rem; width:{table_width}px; }}
    thead th, tbody td {{ overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }}
    thead th {{ font-weight:500; color:{MUTED}; padding:6px 10px; border-bottom:1.5px solid {LINE}; text-align:right; }}
    thead th.mo {{ text-align:center; color:{INK}; }}
    thead th.oc-o {{ color:{GREEN}; }}
    thead th.oc-c {{ color:{VERMILION}; }}
    tbody td {{ padding:6px 10px; border-bottom:1px solid {TABLE_LINE}; text-align:right; }}
    tbody tr:hover td {{ background:{ROW_HOVER}; }}
    td.net-up {{ background:{ACCENT_SOFT}; color:{GREEN_DEEP}; font-weight:600; }}
    td.net-down {{ background:#FBEAEA; color:{VERMILION}; font-weight:600; }}
    td.region, th.region {{
        text-align:left; position:sticky; left:0; background:#FFFFFF; font-weight:500;
        border-right:1.5px solid {LINE};
    }}
    tbody tr:hover td.region {{ background:{ROW_HOVER}; }}
    .bf {{ color:{MUTED}; }}
    tr.region-toggle {{ cursor:pointer; }}
    .chev {{ display:inline-block; width:11px; color:{MUTED}; font-size:.7rem; }}
    td.district-cell {{ padding-left:28px; color:{MUTED}; font-weight:400; }}
    </style></head>
    <body>
    <div class="scroll">
    <table>
        <colgroup>{colgroup_html}</colgroup>
        <thead><tr><th class="region"></th>{header_months}</tr><tr>{header_sub}</tr></thead>
        <tbody id="tbody">{''.join(rows_html)}</tbody>
    </table>
    </div>
    <script>
    document.getElementById('tbody').addEventListener('click', function(e) {{
        var tr = e.target.closest('tr[data-toggle]');
        if (!tr) return;
        var key = tr.getAttribute('data-toggle');
        var open = tr.getAttribute('data-open') === '1';
        document.querySelectorAll('tr[data-group="' + key + '"]').forEach(function(r) {{
            r.style.display = open ? 'none' : '';
        }});
        tr.setAttribute('data-open', open ? '0' : '1');
        var chev = tr.querySelector('.chev');
        if (chev) {{ chev.textContent = open ? '▸' : '▾'; }}
        resizeFrame();
    }});
    function resizeFrame() {{
        try {{
            if (window.frameElement) {{
                var h = document.documentElement.scrollHeight;
                if (h > 0) {{ window.frameElement.style.height = h + 'px'; }}
            }}
        }} catch (err) {{ /* 크로스 오리진 등으로 접근 불가하면 기본 높이를 그대로 쓴다 */ }}
    }}
    resizeFrame();
    window.addEventListener('load', resizeFrame);
    setTimeout(resizeFrame, 300);
    setTimeout(resizeFrame, 1000);
    // st.tabs()의 비활성 탭 안에서 이 iframe이 먼저 만들어지면(숨김 상태) 처음 계산한 높이가
    // 0으로 잡힐 수 있다. 탭을 클릭해 실제로 보이게 되는 순간을 ResizeObserver로 잡아 다시 잰다.
    if (window.ResizeObserver) {{
        new ResizeObserver(resizeFrame).observe(document.body);
    }}
    </script>
    </body></html>
    """


def _render_detail_table(
    table_region_short: str, table_df: pd.DataFrame, regions_full: list, regions_short: list,
    period_months: list, months_desc: list,
) -> None:
    if table_region_short == "전국":
        html = _build_nationwide_tree_component(table_df, regions_full, regions_short, period_months, months_desc)
        components.html(html, height=560, scrolling=False)
    else:
        df, regions = _resolve_table_scope(table_region_short, table_df, regions_full, period_months)
        st.markdown(_build_table_html(df, regions, months_desc), unsafe_allow_html=True)


def _matrix_to_excel_bytes(df: pd.DataFrame, regions: list, months_desc: list) -> bytes:
    """지역(행) x 월(열, 개업/폐업/순증감/영업중) 형태 엑셀. 순증감 셀은 양/음에 따라 배경색을 넣는다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "지역별_상세"

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right")
    metric_labels = ["개업", "폐업", "순증감", "영업중"]

    region_header = ws.cell(row=1, column=1, value="지역")
    region_header.font = bold
    region_header.alignment = center
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)

    col = 2
    for m in months_desc:
        start_col = col
        for label in metric_labels:
            c = ws.cell(row=2, column=col, value=label)
            c.font = bold
            c.alignment = center
            col += 1
        month_cell = ws.cell(row=1, column=start_col, value=m)
        month_cell.font = bold
        month_cell.alignment = center
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=col - 1)

    row = 3
    for region in regions:
        label = _region_label(region)
        ws.cell(row=row, column=1, value=label)
        sub = df[df["지역"] == region].set_index("월") if not df.empty else None
        col = 2
        for m in months_desc:
            opened = closed = live = None
            if sub is not None and m in sub.index:
                o, c, lv = sub.loc[m, "개업"], sub.loc[m, "폐업"], sub.loc[m, "영업중"]
                opened = None if pd.isna(o) else int(o)
                closed = None if pd.isna(c) else int(c)
                live = None if pd.isna(lv) else int(lv)
            net = None if opened is None or closed is None else opened - closed
            for key, value in (("개업", opened), ("폐업", closed), ("순증감", net), ("영업중", live)):
                cell = ws.cell(row=row, column=col, value=value)
                cell.alignment = right
                if key == "순증감" and value is not None:
                    cell.fill = _NET_UP_FILL if value >= 0 else _NET_DOWN_FILL
                col += 1
        row += 1

    ws.freeze_panes = "B3"
    ws.column_dimensions["A"].width = 16
    for c in range(2, col):
        label = metric_labels[(c - 2) % len(metric_labels)]
        ws.column_dimensions[ws.cell(row=2, column=c).column_letter].width = max(9, len(label) + 4)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _leads_to_excel_bytes(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "신규_영업_타겟"
    headers = ["상호명", "법정동", "등록일자"]
    bold = Font(bold=True)
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.font = bold
    for row_idx, r in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=r["상호명"])
        ws.cell(row=row_idx, column=2, value=r["법정동"])
        ws.cell(row=row_idx, column=3, value=r["등록일자"])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 14
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render():
    rows = real_stats.load_region_monthly_stats()
    if not rows:
        st.info("표시할 데이터가 없습니다.")
        return

    df = pd.DataFrame(rows)
    all_months = sorted(df["월"].unique())
    latest_month = all_months[-1]
    regions_short = ru.ordered_short_regions(df["지역"].unique())
    regions_full = [ru.FULL_NAME[s] for s in regions_short]
    snapshot_date = real_stats.load_latest_snapshot_date()

    with st.container(key="brokers_root"):
        st.markdown(_FILTER_CSS, unsafe_allow_html=True)
        st.markdown(
            f'<div class="ov-topline">'
            f'<div class="sub">스냅샷 기준 {snapshot_date or latest_month}</div>'
            f"</div>",
            unsafe_allow_html=True,
        )

        # ---- 공통 필터 상태 기본값 (KPI 카드가 필터보다 먼저 그려지므로 위젯 만들기 전에 session_state부터 채운다) ----
        st.session_state.setdefault("br_region", regions_short[0])
        st.session_state.setdefault("br_start", all_months[0])
        st.session_state.setdefault("br_end", latest_month)
        st.session_state.setdefault("br_preset", "전체")
        st.session_state.setdefault("br_table_region", "전국")

        region_short = st.session_state["br_region"]
        region = ru.FULL_NAME[region_short]
        start_month, end_month = st.session_state["br_start"], st.session_state["br_end"]
        if start_month > end_month:
            start_month, end_month = end_month, start_month
        period_months = [m for m in all_months if start_month <= m <= end_month]
        months_desc = list(reversed(period_months))

        region_series = sorted([dict(r) for r in rows if r["지역"] == region], key=lambda r: r["월"])
        plot_rows = [r for r in region_series if r["월"] in period_months]
        pdf = pd.DataFrame(plot_rows)

        # ---- KPI 1행 (현재 스냅샷 기준 + 선택된 지역·기간 기준) ----
        status_counts = real_stats.load_office_status_counts()
        active_cnt = status_counts.get("영업중", 0)

        new_offices = real_stats.load_new_offices_this_month()
        national = real_stats.national_monthly_series()
        latest_national = national[-1] if national else {}

        # 상단 지역·기간 필터가 기본값('전체' 기간)이면 이달 기준 개업/폐업을, 기간을 좁히면(프리셋 포함)
        # 선택된 지역·기간 동안의 누적 개업/폐업을 보여준다. '전체' 기간 누적은 이 카드의 취지(당월
        # 펄스체크)에 안 맞아서 그 경우엔 이달 기준으로 되돌린다.
        is_full_range = (start_month, end_month) == (all_months[0], latest_month)
        if is_full_range:
            period_rows = [r for r in plot_rows if r["월"] == latest_month]
            title_opened, title_closed, title_net = "이달의 신규 개업", "당월 폐업", "당월 순증감"
            sub_opened = f'{snapshot_date or latest_month} 기준'
            sub_closed = sub_net = f'{latest_month} 기준'
        else:
            period_rows = plot_rows
            range_txt = start_month if start_month == end_month else f'{start_month}~{end_month}'
            title_opened, title_closed, title_net = "기간 내 신규 개업", "기간 내 폐업", "기간 내 순증감"
            sub_opened = sub_closed = sub_net = f'{range_txt} · {region_short}'
        opened_period = sum(r.get("개업") or 0 for r in period_rows)
        closed_period = sum(r.get("폐업") or 0 for r in period_rows)
        net_period = opened_period - closed_period
        net_cls = "plus" if net_period >= 0 else "minus"

        cards_html = (
            '<div class="ov-stats">'
            + '<div class="ov-stat open"><div class="label">실제 영업 중 사무소</div>'
              f'<div class="value">{active_cnt:,}</div><div class="delta">전체 유효 등록 건수</div></div>'
            + f'<div class="ov-stat open"><div class="label">{title_opened}</div>'
              f'<div class="value">{opened_period:,.0f}</div><div class="delta">{sub_opened}</div></div>'
            + f'<div class="ov-stat close"><div class="label">{title_closed}</div>'
              f'<div class="value">{closed_period:,.0f}</div><div class="delta">{sub_closed}</div></div>'
            + f'<div class="ov-stat net"><div class="label">{title_net}</div>'
              f'<div class="value {net_cls}">{net_period:+,.0f}</div><div class="delta">{sub_net}</div></div>'
            + "</div>"
        )
        st.markdown(cards_html, unsafe_allow_html=True)

        # ---- 공통 필터 위젯: 지역 + 기간(기본 전체) — 차트/매트릭스 표가 함께 따른다 ----
        # 지역/시작/종료 + 빠른선택 pill을 전부 한 줄에 컴팩트하게(_FILTER_CSS 참고).
        with st.container(key="br_filter_row"):
            c1, c2, c3, c4 = st.columns([1, 0.85, 0.85, 2.6])
            region_short = c1.selectbox("지역", regions_short, key="br_region", on_change=_sync_table_region)
            region = ru.FULL_NAME[region_short]
            start_month = c2.selectbox("시작", all_months, key="br_start")
            end_month = c3.selectbox("종료", all_months, key="br_end")
            with c4:
                c4.markdown('<div style="height:1.6rem"></div>', unsafe_allow_html=True)
                with st.container(key="br_preset_row"):
                    st.segmented_control(
                        "기간 프리셋",
                        options=["최근 6개월", "올해", "3개년", "5개년", "전체"],
                        label_visibility="collapsed",
                        key="br_preset",
                        on_change=_apply_preset,
                        args=(all_months, latest_month),
                    )

        if start_month > end_month:
            start_month, end_month = end_month, start_month

        # ---- 탭 3개: 그래프 추이 / 지역별 상세 / B2B 신규 영업 타겟 ----
        tab1, tab2, tab3 = st.tabs(["그래프 추이", "지역별 상세", "B2B 신규 영업 타겟"])

        with tab1:
            with st.container(border=True):
                st.markdown(f'<div class="ov-panel-title">{region_short} 개폐업 추이</div>', unsafe_allow_html=True)
                st.markdown('<div class="ov-panel-desc">초록(개업)·주홍(폐업) 막대, 파란선(순증감)</div>', unsafe_allow_html=True)
                fig1 = go.Figure()
                fig1.add_bar(x=pdf["월"], y=pdf["개업"], name="개업", marker_color=GREEN)
                fig1.add_bar(x=pdf["월"], y=pdf["폐업"], name="폐업", marker_color=VERMILION)
                fig1.add_trace(go.Scatter(
                    x=pdf["월"], y=pdf["개업"] - pdf["폐업"], name="순증감",
                    mode="lines+markers", line=dict(color=BLUE, width=2),
                ))
                fig1.update_layout(
                    barmode="group",
                    plot_bgcolor=CARD, paper_bgcolor=CARD,
                    font=dict(family="Pretendard Variable, Pretendard, sans-serif", color=INK, size=12),
                    yaxis=dict(title="건수", gridcolor=LINE, zerolinecolor=LINE, tickformat=","),
                    xaxis=dict(gridcolor=LINE),
                    legend=dict(orientation="h", yanchor="bottom", y=1.0, x=0),
                    hovermode="x unified",
                    margin=dict(l=10, r=10, t=10, b=10),
                    height=340,
                    dragmode=False,
                )
                st.plotly_chart(fig1, use_container_width=True, config={"displayModeBar": False})

        with tab2:
            with st.container(border=True):
                h1, hl, h2, h_actions = st.columns([2.5, 0.35, 0.65, 1.2])
                h1.markdown('<div class="ov-panel-title" style="padding-top:.4rem;">지역별 상세 매트릭스 표</div>', unsafe_allow_html=True)
                hl.markdown('<div style="padding-top:.5rem; font-size:.83rem; color:#6B756E;">지역</div>', unsafe_allow_html=True)
                table_region_short = h2.selectbox("지역", regions_short, key="br_table_region", label_visibility="collapsed")

                table_df = df[df["월"].isin(period_months)][["월", "지역", "개업", "폐업", "영업중", "출처"]]
                df_matrix, table_regions = _resolve_table_scope(table_region_short, table_df, regions_full, period_months)
                excel_bytes = _matrix_to_excel_bytes(df_matrix, table_regions, months_desc)
                h3, hr = h_actions.columns([0.78, 0.22])
                h3.download_button(
                    "⬇ 엑셀 다운로드",
                    data=excel_bytes,
                    file_name=f"지역별_상세매트릭스_{start_month}_{end_month}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_matrix",
                    use_container_width=True,
                )
                hr.button("↺", key="reset_matrix_region", on_click=_reset_table_region, help="전국으로 되돌리기", use_container_width=True)
                st.caption("※ 순증감 셀은 양수면 초록, 음수면 빨강 배경으로 표시됩니다. 시군구·영업중(Live) 데이터는 2026-09월부터 제공됩니다.")
                _render_detail_table(table_region_short, table_df, regions_full, regions_short, period_months, months_desc)

        with tab3:
            with st.container(border=True):
                top1, top2 = st.columns([3, 1])
                top1.markdown(
                    f'<div class="ov-panel-title">이달({latest_national.get("월", latest_month)})의 신규 개업 리스트</div>'
                    '<div class="ov-panel-desc">등록일자가 이번 달인 사무소입니다. (전화번호는 현재 수집 데이터에 없어 표시되지 않습니다)</div>',
                    unsafe_allow_html=True,
                )
                if new_offices:
                    excel_bytes = _leads_to_excel_bytes(new_offices)
                    top2.markdown('<div style="height:.3rem"></div>', unsafe_allow_html=True)
                    top2.download_button(
                        "⬇ 엑셀 다운로드",
                        data=excel_bytes,
                        file_name=f"B2B_신규영업타겟_{latest_national.get('월', latest_month)}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_leads",
                        use_container_width=True,
                    )
                    lead_df = pd.DataFrame(new_offices)
                    st.dataframe(lead_df, use_container_width=True, hide_index=True)
                else:
                    st.info("이번 달 등록된 신규 개업 사무소가 없습니다.")
