# -*- coding: utf-8 -*-
"""
매일 수집 결과를 엑셀에 쌓는다.
- 행: 지역(시군구), 맨 위에 "전체"(전국 합계) 행
- 열: 날짜가 오른쪽으로 계속 늘어남 (날짜 하나당 영업중/신규등록/폐업 3칸)
같은 날짜로 다시 실행하면 그 날짜의 열을 덮어쓴다(중복 추가 안 됨).
"""
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

import region_utils as ru

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, "data", "지역별_일별_통계.xlsx")
SHEET_NAME = "일별_지역별_통계"
METRICS = ["영업중", "신규등록", "폐업"]
TOTAL_LABEL = "전체(전국)"

HEADER_ROW = 1
SUBHEADER_ROW = 2
TOTAL_ROW = 3
FIRST_REGION_ROW = 4


def _region_sort_key(region: str):
    sido_short = ru.SHORT_NAME.get(ru.normalize_sido(region), None)
    order = ru.REGION_ORDER.index(sido_short) if sido_short in ru.REGION_ORDER else len(ru.REGION_ORDER)
    return (order, region)


def _new_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.cell(row=HEADER_ROW, column=1, value="지역")
    ws.cell(row=TOTAL_ROW, column=1, value=TOTAL_LABEL).font = Font(bold=True)
    ws.freeze_panes = f"B{FIRST_REGION_ROW}"
    ws.column_dimensions["A"].width = 24
    return wb, ws


def _find_date_columns(ws, date_str: str):
    for col in range(2, ws.max_column + 1, len(METRICS)):
        if ws.cell(row=HEADER_ROW, column=col).value == date_str:
            return col
    return None


def _find_or_add_region_row(ws, region: str) -> int:
    for row in range(FIRST_REGION_ROW, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == region:
            return row
    new_row = max(ws.max_row + 1, FIRST_REGION_ROW)
    ws.cell(row=new_row, column=1, value=region)
    return new_row


def append_daily_stats(today: str, active_by_region: dict, opened_by_region: dict, closed_by_region: dict) -> str:
    if os.path.exists(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    else:
        wb, ws = _new_workbook()

    start_col = _find_date_columns(ws, today)
    if start_col is None:
        start_col = ws.max_column + 1 if ws.max_column > 1 else 2
        ws.merge_cells(start_row=HEADER_ROW, start_column=start_col, end_row=HEADER_ROW, end_column=start_col + len(METRICS) - 1)
        header_cell = ws.cell(row=HEADER_ROW, column=start_col, value=today)
        header_cell.alignment = Alignment(horizontal="center")
        header_cell.font = Font(bold=True)
        for i, metric in enumerate(METRICS):
            ws.cell(row=SUBHEADER_ROW, column=start_col + i, value=metric)

    regions = sorted(set(active_by_region) | set(opened_by_region) | set(closed_by_region), key=_region_sort_key)

    ws.cell(row=TOTAL_ROW, column=start_col, value=sum(active_by_region.values()))
    ws.cell(row=TOTAL_ROW, column=start_col + 1, value=sum(opened_by_region.values()))
    ws.cell(row=TOTAL_ROW, column=start_col + 2, value=sum(closed_by_region.values()))

    for region in regions:
        row = _find_or_add_region_row(ws, region)
        ws.cell(row=row, column=start_col, value=active_by_region.get(region, 0))
        ws.cell(row=row, column=start_col + 1, value=opened_by_region.get(region, 0))
        ws.cell(row=row, column=start_col + 2, value=closed_by_region.get(region, 0))

    os.makedirs(os.path.join(BASE_DIR, "data"), exist_ok=True)
    wb.save(EXCEL_FILE)
    return EXCEL_FILE


def rebuild_from_db():
    """data/offices.db 의 daily_region_stats 전체를 가지고 엑셀을 처음부터 새로 만든다."""
    import db

    if os.path.exists(EXCEL_FILE):
        os.remove(EXCEL_FILE)

    conn = db.get_conn()
    dates = [r["date"] for r in conn.execute("SELECT DISTINCT date FROM daily_region_stats ORDER BY date")]
    for d in dates:
        rows = conn.execute(
            "SELECT region, active, opened, closed FROM daily_region_stats WHERE date=?", (d,)
        ).fetchall()
        active = {r["region"]: r["active"] or 0 for r in rows}
        opened = {r["region"]: r["opened"] or 0 for r in rows}
        closed = {r["region"]: r["closed"] or 0 for r in rows}
        append_daily_stats(d, active, opened, closed)
    conn.close()
    return EXCEL_FILE


if __name__ == "__main__":
    path = rebuild_from_db()
    print(f"재생성 완료: {path}")
